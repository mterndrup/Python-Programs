import os
import openpyxl

# Directories
base_dir = r"C:\Users\ketgl\OneDrive\Desktop\Sandbox\Cactus-All"
output_dir = r"C:\Users\ketgl\OneDrive\Desktop\Sandbox\Logs"
output_file = os.path.join(output_dir, "ITSx_counts.xlsx")

# Regions we care about (ITSx outputs)
expected_regions = ["5.8S", "ITS1", "ITS2", "LSU", "SSU", "Full"]

# Cutadapt primer outputs
cutadapt_regions = ["18S", "FITS", "PITS"]

# Map raw ITSx suffix to expected region names
region_map = {
    "5_8S": "5.8S",
    "ITS1": "ITS1",
    "ITS2": "ITS2",
    "LSU": "LSU",
    "SSU": "SSU",
    "full": "Full"
}

def count_fasta_sequences(fasta_path):
    """Count sequences in a fasta file by number of '>' lines."""
    count = 0
    try:
        with open(fasta_path, "r") as f:
            for line in f:
                if line.startswith(">"):
                    count += 1
    except FileNotFoundError:
        count = 0
    return count

def process_itsx_folder(folder_path, barcode_folder):
    """Process ITSx folder (All or PITS)."""
    row = {"Barcode": barcode_folder}
    total = 0
    for region in expected_regions:
        row[region] = 0

    if os.path.isdir(folder_path):
        for fname in os.listdir(folder_path):
            if not fname.endswith(".fasta"):
                continue
            if "_ITSx." in fname:
                region_part = fname.split("_ITSx.")[-1].replace(".fasta", "")
                if region_part in region_map:
                    region = region_map[region_part]
                    fasta_path = os.path.join(folder_path, fname)
                    count = count_fasta_sequences(fasta_path)
                    row[region] = count

    # compute ITSx total (exclude Full)
    for region in expected_regions:
        if region != "Full":
            total += row[region]
    row["Total"] = total
    return row

# --- Collect results ---
results_all = []
results_pits = []
results_full = []  # All + PITS + Diff
grand_totals = {region: 0 for region in expected_regions}
grand_totals["Total"] = 0

for barcode_folder in os.listdir(base_dir):
    barcode_path = os.path.join(base_dir, barcode_folder)
    if not os.path.isdir(barcode_path):
        continue

    # ITSx-All
    itsx_all_dir = os.path.join(barcode_path, "ITSx-All")
    row_all = process_itsx_folder(itsx_all_dir, barcode_folder)

    # Cutadapt counts
    cutadapt_total = 0
    for primer in cutadapt_regions:
        primer_file = os.path.join(barcode_path, f"{barcode_folder}_{primer}.fasta")
        row_all[primer] = 0
        if os.path.exists(primer_file):
            row_all[primer] = count_fasta_sequences(primer_file)
        cutadapt_total += row_all[primer]
    row_all["CutadaptTotal"] = cutadapt_total

    results_all.append(row_all)

    # update ITSx grand totals (not cutadapt)
    for region in expected_regions:
        grand_totals[region] += row_all[region]
    grand_totals["Total"] += row_all["Total"]

    # ITSx-PITS
    itsx_pits_dir = os.path.join(barcode_path, "ITSx-PITS")
    row_pits = process_itsx_folder(itsx_pits_dir, barcode_folder)
    results_pits.append(row_pits)

    # Full comparison row
    row_full = {"Barcode": barcode_folder}
    for region in expected_regions + ["Total"]:
        row_full[f"{region} (All)"] = row_all.get(region, 0)
        row_full[f"{region} (PITS)"] = row_pits.get(region, 0)
        row_full[f"{region} (Diff)"] = row_all.get(region, 0) - row_pits.get(region, 0)
    results_full.append(row_full)

# --- Write Excel workbook with three sheets ---
os.makedirs(output_dir, exist_ok=True)

wb = openpyxl.Workbook()

# Sheet1: ITSx-All + Cutadapt
ws1 = wb.active
ws1.title = "ITSx-All + Cutadapt"
headers1 = ["Barcode"] + expected_regions + ["Total"] + cutadapt_regions + ["CutadaptTotal"]
ws1.append(headers1)
for row in results_all:
    ws1.append([row.get(h, 0) for h in headers1])

# Sheet2: ITSx-PITS raw counts
ws2 = wb.create_sheet("ITSx-PITS")
headers2 = ["Barcode"] + expected_regions + ["Total"]
ws2.append(headers2)
for row in results_pits:
    ws2.append([row.get(h, 0) for h in headers2])

# Sheet3: Full Comparison (All vs PITS vs Diff)
ws3 = wb.create_sheet("All vs PITS vs Diff")
headers3 = ["Barcode"]
for region in expected_regions + ["Total"]:
    headers3 += [f"{region} (All)", f"{region} (PITS)", f"{region} (Diff)"]
ws3.append(headers3)
for row in results_full:
    ws3.append([row.get(h, 0) for h in headers3])

wb.save(output_file)

print(f"✅ Done! Results saved to {output_file}\n")

# Print grand totals to screen (ITSx-All only)
print("=== Grand Totals Across All Barcodes (ITSx-All only) ===")
for region in expected_regions + ["Total"]:
    print(f"{region}: {grand_totals[region]}")
